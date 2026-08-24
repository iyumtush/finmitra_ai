import collections
import collections.abc
collections.Container = collections.abc.Container
collections.Iterable = collections.abc.Iterable
collections.Sequence = collections.abc.Sequence
collections.Mapping = collections.abc.Mapping
collections.MutableMapping = collections.abc.MutableMapping
collections.MutableSequence = collections.abc.MutableSequence
collections.MutableSet = collections.abc.MutableSet
collections.Callable = collections.abc.Callable

from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]

images = [
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/seminar_slide_1.png',
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/seminar_slide_2.png',
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/seminar_slide_3.png',
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/seminar_slide_4.png',
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/seminar_slide_5.png',
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/seminar_slide_6.png'
]

for img_path in images:
    slide = prs.slides.add_slide(blank_slide_layout)
    slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

out_path = '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/FinMitra_Seminar_Presentation.pptx'
prs.save(out_path)
print("Saved PPTX to", out_path)
