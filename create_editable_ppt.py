import collections
import collections.abc
collections.Container = collections.abc.Container
collections.Iterable = collections.abc.Iterable
collections.Sequence = collections.abc.Sequence
collections.Mapping = collections.abc.Mapping
collections.MutableMapping = collections.abc.MutableMapping
collections.MutableSequence = collections.abc.MutableSequence
collections.MutableSet = collections.abc.MutableSet
collections.Callable = collections.abc.Callable

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

title_slide_layout = prs.slide_layouts[0]
title_and_content_layout = prs.slide_layouts[1]
title_only_layout = prs.slide_layouts[5]

# Slide 1: Project Aim & Objectives
slide1 = prs.slides.add_slide(title_and_content_layout)
slide1.shapes.title.text = "Project Aim & Objectives"
tf1 = slide1.placeholders[1].text_frame
tf1.text = "To develop an intelligent, context-aware personal financial management platform (FinMitra) that automates expense tracking, simplifies budgeting, and provides highly personalized financial guidance using AI."
p = tf1.add_paragraph()
p.text = "1. Smart Financial Mgmt: Centralized dashboard for tracking income & expenses."
p = tf1.add_paragraph()
p.text = "2. Personalized Guidance: Context-aware recommendations based on live data."
p = tf1.add_paragraph()
p.text = "3. Financial Literacy: Educational assistance via an AI conversational agent."
p = tf1.add_paragraph()
p.text = "4. Better Decision-Making: Budget tracking and spending analysis."
p = tf1.add_paragraph()
p.text = "5. Security & Privacy: Secure authentication and Row Level Security."
p = tf1.add_paragraph()
p.text = "6. AI Automation: Automated multimodal receipt parsing."

# Slide 2: Work Done Till Previous Seminar
slide2 = prs.slides.add_slide(title_and_content_layout)
slide2.shapes.title.text = "Work Done Till Previous Seminar"
tf2 = slide2.placeholders[1].text_frame
tf2.text = "Features Developed:"
tf2.add_paragraph().text = "• Secure Authentication & Dashboard: Real-time summary dashboard for Income, Expenses, and Savings."
tf2.add_paragraph().text = "• Transaction Management: Full CRUD operations with smart categorization."
tf2.add_paragraph().text = "• Goal-Oriented Budgeting: Custom category budgets with visual progress indicators."
tf2.add_paragraph().text = "• Context-Aware AI Chatbot: Gemini-powered assistant answering questions on live data."
tf2.add_paragraph().text = "• Multimodal Receipt Parsing: Image uploads for automated transaction entry."
tf2.add_paragraph().text = "• Persistent Chat History: Dedicated history view for past AI conversations."

# Slide 3: Platform Showcase
slide3 = prs.slides.add_slide(title_only_layout)
slide3.shapes.title.text = "Platform Showcase & Current Progress"
images = [
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/.user_uploaded/media_1787593246741.png', # Auth
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/.user_uploaded/media_1787593221101.png', # Dashboard
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/.user_uploaded/media_1787593332893.png', # Budgets
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/.user_uploaded/media_1787593394644.png', # AI Insights
    '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/.user_uploaded/media_1787593271715.png'  # Transactions
]
captions = ["Secure Auth", "Dashboard", "Budgets", "AI Insights", "Transactions"]

# Top row (3 images)
for i in range(3):
    left = Inches(0.5 + i * 4.2)
    top = Inches(1.5)
    width = Inches(4.0)
    height = Inches(2.4)
    slide3.shapes.add_picture(images[i], left, top, width=width, height=height)
    txBox = slide3.shapes.add_textbox(left, top + height, width, Inches(0.5))
    txBox.text_frame.text = captions[i]

# Bottom row (2 images)
for i in range(3, 5):
    idx = i - 3
    left = Inches(1.5 + idx * 5.5)
    top = Inches(4.3)
    width = Inches(5.0)
    height = Inches(2.7)
    slide3.shapes.add_picture(images[i], left, top, width=width, height=height)
    txBox = slide3.shapes.add_textbox(left, top + height, width, Inches(0.5))
    txBox.text_frame.text = captions[i]

# Slide 4: Result & Discussion
slide4 = prs.slides.add_slide(title_and_content_layout)
slide4.shapes.title.text = "Result & Discussion"
tf4 = slide4.placeholders[1].text_frame
tf4.text = "• High Accuracy in Parsing: Gemini Vision accurately extracts transaction amounts, dates, and categories from varied receipt formats, eliminating tedious manual entry."
tf4.add_paragraph().text = "• Contextual AI Awareness: Unlike generic AI bots, our system successfully parses the user's specific live database (income/expenses/budgets) to provide actionable, tailored advice."
tf4.add_paragraph().text = "• Enhanced User Experience: The seamless transition between active chat and historical logs provides a fluid, robust personal wealth management experience."

# Slide 5: Conclusion
slide5 = prs.slides.add_slide(title_and_content_layout)
slide5.shapes.title.text = "Conclusion (Of Current Work)"
tf5 = slide5.placeholders[1].text_frame
tf5.text = "• The integration of the Google Gemini Multimodal AI has successfully transformed FinMitra from a static expense tracker into a proactive financial assistant."
tf5.add_paragraph().text = "• Our multimodal capabilities reliably automate data entry through image parsing, significantly reducing user friction."
tf5.add_paragraph().text = "• Simultaneously, the custom context-engine bridges the gap between raw database metrics and human-readable guidance."
tf5.add_paragraph().text = "• FinMitra is now a secure, robust, and highly intelligent foundation, fully prepared for future additions like long-term investment planning and EMI affordability analysis."

# Slide 6: Achievement of Objective Number
slide6 = prs.slides.add_slide(title_and_content_layout)
slide6.shapes.title.text = "Achievement of Objective Number"
tf6 = slide6.placeholders[1].text_frame
tf6.text = "1. Smart Financial Mgmt: ✓ Achieved (Users efficiently manage income, expenses, and budgets inside one intelligent dashboard.)"
tf6.add_paragraph().text = "2. Personalized Guidance: ✓ Achieved (The Chatbot analyzes user behavior and provides tailored savings recommendations based on actual records.)"
tf6.add_paragraph().text = "3. Improved Financial Literacy: ↻ In Progress (The Chatbot answers financial queries correctly, establishing awareness and foundational education.)"
tf6.add_paragraph().text = "4. Better Decision-Making: ↻ In Progress (Budget tracking gives a good start, but deeper AI analysis on loan/EMI affordability is still under development.)"
tf6.add_paragraph().text = "5. Security and Privacy: ✓ Achieved (Sensitive financial data is fully protected using Supabase Authentication and Row Level Security.)"
tf6.add_paragraph().text = "6. Intelligent AI Automation: ✓ Achieved (Integrated Google Gemini Multimodal AI parses receipt images automatically, removing manual data entry.)"

out_path = '/Users/admin/.gemini/antigravity-ide/brain/a6016e96-3b3e-4478-a934-f24070502ee9/FinMitra_Seminar_Presentation.pptx'
prs.save(out_path)
print("Saved fully editable PPTX to", out_path)
